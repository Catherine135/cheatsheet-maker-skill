"""Build a source-derived editable Word cheatsheet from semantically marked runs.

Usage: python build_cheatsheet.py content.json output.docx [--profile approved]
The JSON is task-local scratch data, not a user-facing deliverable.
"""
from pathlib import Path
from copy import deepcopy
from functools import lru_cache
import argparse,json,re
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
SOURCES={'approved':ROOT/'assets/reference.docx'}
ROLES={'body','major','topic','term','concept','rule','code','code_core','identifier','identifier_core','comment','output','gloss','circle'}
COLORS={'major':('yellow','000000'),'topic':('cyan','000000'),
        'concept':('blue','FFFFFF'),'code':('black','FFFFFF'),'code_core':('black','95DCF7'),
        'comment':('lightGray','000000'),'output':('lightGray','000000'),
        'identifier':('lightGray','000000'),'identifier_core':('lightGray','95DCF7'),
        'rule':(None,'EE0000')}
CJK=r'[\u3400-\u9fff\u3000-\u303f\uff01-\uff60\uffe0-\uffe6]'

def prop(parent,name,**attributes):
    el=parent.find(qn('w:'+name))
    if el is None:el=OxmlElement('w:'+name);parent.append(el)
    el.attrib.clear()
    for k,v in attributes.items():el.set(qn('w:'+k),str(v))
    return el

@lru_cache(maxsize=1)
def circle_template():
    """Use the approved Word GB3 component, not a font approximation."""
    reference=Document(SOURCES['approved'])
    sequences=[];current=None
    for p in reference.paragraphs:
        for run in p._p.findall(qn('w:r')):
            field=run.find(qn('w:fldChar'))
            kind=field.get(qn('w:fldCharType')) if field is not None else None
            if kind=='begin':current=[]
            if current is not None:current.append(deepcopy(run))
            if kind=='end' and current is not None:
                instr=''.join((n.text or '') for r in current for n in r.findall(qn('w:instrText')))
                if re.fullmatch(r'\s*=\s*1\s+\\\*\s+GB3\s*',instr):sequences.append(current)
                current=None
    if not sequences:raise ValueError('Approved reference has no circle-one GB3 field')
    return sequences[-1]

def add_circle(p,char):
    sequence=deepcopy(circle_template());number=ord(char)-ord('①')+1
    for r in sequence:
        for instruction in r.findall(qn('w:instrText')):
            if 'GB3' in (instruction.text or ''):instruction.text=f'= {number} \\* GB3'
        for result in r.findall(qn('w:t')):result.text=char
        p._p.append(r)

def paragraph_mark(p):
    """Match font metrics without copying visible highlighting to the mark."""
    pp=p._p.get_or_add_pPr();mark=prop(pp,'rPr')
    last=None
    for r in p._p.findall(qn('w:r')):
        if r.find(qn('w:t')) is not None:last=r.find(qn('w:rPr'))
    for tag in ('rFonts','b','bCs','sz','szCs'):
        src=last.find(qn('w:'+tag)) if last is not None else None
        if src is not None:mark.append(deepcopy(src))
    if mark.find(qn('w:rFonts')) is None:
        prop(mark,'rFonts',ascii='Abadi MT Condensed Light',hAnsi='Abadi MT Condensed Light',eastAsiaTheme='majorHAnsi')
    for tag in ('sz','szCs'):
        if mark.find(qn('w:'+tag)) is None:prop(mark,tag,val=13)

def normalize_defaults(d):
    styles=d.styles.element
    defaults=styles.find(qn('w:docDefaults'))
    if defaults is None:defaults=OxmlElement('w:docDefaults');styles.insert(0,defaults)
    rd=defaults.find(qn('w:rPrDefault'))
    if rd is None:rd=OxmlElement('w:rPrDefault');defaults.append(rd)
    rp=rd.find(qn('w:rPr'))
    if rp is None:rp=OxmlElement('w:rPr');rd.append(rp)
    for target in (rp,d.styles['Normal'].element.get_or_add_rPr()):
        prop(target,'rFonts',ascii='Abadi MT Condensed Light',hAnsi='Abadi MT Condensed Light',eastAsiaTheme='majorHAnsi')
        prop(target,'sz',val=13);prop(target,'szCs',val=13)
        prop(target,'b',val=0);prop(target,'bCs',val=0)
    normal=d.styles['Normal'].paragraph_format
    normal.space_before=Pt(0);normal.space_after=Pt(0)
    normal.line_spacing=Pt(6);normal.line_spacing_rule=WD_LINE_SPACING.EXACTLY

def add_text(p,text,role,profile,bold=None,size_pt=None):
    if role not in ROLES:raise ValueError('Unknown run role: '+role)
    # Split circled numerals separately even inside a bold, colored or code run.
    # Split Han text to reproduce the deliberate 5.5 pt Chinese gloss system.
    chunks=re.findall(r'[①-⑳]|'+CJK+r'+|[^①-⑳\u3400-\u9fff\u3000-\u303f\uff01-\uff60\uffe0-\uffe6]+',text)
    for chunk in chunks:
        circ=bool(re.fullmatch(r'[①-⑳]',chunk))
        if circ:
            add_circle(p,chunk);continue
        han=bool(re.fullmatch(CJK+r'+',chunk))
        size=7.5 if role=='major' else 5.5 if circ or role in {'circle','gloss'} or han else 6.5
        if size_pt is not None and not circ:size=float(size_pt)
        is_bold=role in {'major','topic','term','concept','rule','code','code_core'}
        if bold is not None:is_bold=bool(bold)
        if han:is_bold=True
        # The +1 pt title rule applies to each language's own baseline:
        # English 6.5 -> 7.5 pt, Chinese gloss 5.5 -> 6.5 pt.
        if role=='major' and not circ:size=6.5 if han else 7.5
        r=p.add_run(chunk);rp=r._element.get_or_add_rPr()
        prop(rp,'rFonts',ascii='Abadi MT Condensed Light',hAnsi='Abadi MT Condensed Light',eastAsiaTheme='majorHAnsi')
        prop(rp,'sz',val=round(size*2));prop(rp,'szCs',val=round(size*2))
        prop(rp,'b',val='1' if is_bold else '0');prop(rp,'bCs',val='1' if is_bold else '0')
        highlight,color=COLORS.get(role,(None,'000000'))
        prop(rp,'color',val=color)
        if highlight:prop(rp,'highlight',val=highlight)
    return p

def expand_templates(items):
    """Expand explicit semantic placeholders; never infer identifiers from spelling."""
    for item in items:
        if 'code_template' not in item:
            yield item;continue
        if any(k in item for k in ('text','runs','image','clone_paragraph')):
            raise ValueError('code_template cannot also contain text/runs/image/clone_paragraph')
        lines=item['code_template'].splitlines()
        for i,line in enumerate(lines):
            runs=[]
            for token in re.split(r'(\{\{.*?\}\})',line):
                if not token:continue
                variable=token.startswith('{{') and token.endswith('}}')
                runs.append({'text':token[2:-2] if variable else token,
                             'role':'identifier' if variable else 'code'})
            yield {**{k:v for k,v in item.items() if k!='code_template'},
                   'role':'code','runs':runs,'keep_lines':True,'keep_next':i<len(lines)-1,
                   'column_break_before':i==0 and item.get('column_break_before',False)}

def add_image(d,item):
    image=item['image'];width=float(image.get('width_pt',184))
    sec=d.sections[0];cols=sec._sectPr.find(qn('w:cols'))
    num=int(cols.get(qn('w:num'),'1'));gap=int(cols.get(qn('w:space'),'0'))/20
    column=(sec.page_width.pt-sec.left_margin.pt-sec.right_margin.pt-gap*(num-1))/num
    if width<=0 or width>column:raise ValueError(f'Image width {width} exceeds column {column:.2f}')
    p=d.add_paragraph();p.style='a';pf=p.paragraph_format
    pf.space_before=Pt(0);pf.space_after=Pt(0);pf.line_spacing=1
    pf.left_indent=Pt(0);pf.right_indent=Pt(0);pf.first_line_indent=Pt(0)
    pf.widow_control=False;pf.keep_together=True;pf.keep_with_next=bool(item.get('keep_next',False))
    p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    pp=p._p.get_or_add_pPr();prop(pp,'snapToGrid',val=0)
    mark=prop(pp,'rPr');prop(mark,'sz',val=2);prop(mark,'szCs',val=2)
    r=p.add_run();r.font.size=Pt(1);prop(r._r.get_or_add_rPr(),'szCs',val=2)
    prop(r._r.get_or_add_rPr(),'spacing',val=0)
    shape=r.add_picture(str(Path(image['path']).expanduser()),width=Pt(width))
    for key in ('distT','distB','distL','distR'):shape._inline.set(key,'0')
    if image.get('alt'):shape._inline.docPr.set('descr',image['alt'])
    return p

def build(data,output,profile='approved'):
    if profile not in SOURCES:raise ValueError('Unknown profile: '+profile)
    reference=SOURCES[profile]
    if Path(output).resolve() in [p.resolve() for p in SOURCES.values()]:raise ValueError('Cannot overwrite a retained reference')
    d=Document(reference)
    normalize_defaults(d)
    body=d._element.body
    for child in list(body):
        if child.tag!=qn('w:sectPr'):body.remove(child)
    original=Document(reference)
    chapters=set()
    for item in expand_templates(data['paragraphs']):
        if item.get('column_break_before') and d.paragraphs:
            # Do not put a column break in an auto-height image paragraph:
            # Word/LibreOffice can move the entire picture to the next column.
            bp=d.add_paragraph();bp.style='a';bf=bp.paragraph_format
            bf.space_before=Pt(0);bf.space_after=Pt(0)
            bf.line_spacing=Pt(6);bf.line_spacing_rule=WD_LINE_SPACING.EXACTLY
            bf.widow_control=False;bf.keep_with_next=False;bf.keep_together=False
            bf.left_indent=Pt(0);bf.right_indent=Pt(0);bf.first_line_indent=Pt(0)
            br=OxmlElement('w:br');br.set(qn('w:type'),'column')
            rr=bp.add_run();rr._r.append(br)
            rp=rr._r.get_or_add_rPr()
            prop(rp,'rFonts',ascii='Abadi MT Condensed Light',hAnsi='Abadi MT Condensed Light',eastAsiaTheme='majorHAnsi')
            prop(rp,'sz',val=13);prop(rp,'szCs',val=13)
            paragraph_mark(bp)
        if 'image' in item:
            add_image(d,item);continue
        if 'clone_paragraph' in item:
            # For a regression fixture or intentional reuse of source diagrams/formulas.
            # Original relationship-backed content is valid because the source package is retained.
            clone=deepcopy(original.paragraphs[int(item['clone_paragraph'])]._p)
            body.insert(len(body)-1,clone)
            continue
        role=item.get('role','body')
        if role not in ROLES:raise ValueError('Unknown paragraph role: '+role)
        if role=='major':
            heading=item.get('text',''.join(r['text'] for r in item.get('runs',[])))
            key=str(item.get('chapter_id',heading)).strip().casefold()
            if key in chapters:raise ValueError('Repeated chapter title: '+key)
            chapters.add(key)
        p=d.add_paragraph();p.style='a'
        pf=p.paragraph_format
        pf.space_before=Pt(0);pf.space_after=Pt(0)
        pf.line_spacing=Pt(6);pf.line_spacing_rule=WD_LINE_SPACING.EXACTLY
        pf.widow_control=False
        pf.keep_with_next=bool(item.get('keep_next',role in {'major','topic'}))
        pf.keep_together=bool(item.get('keep_lines',False))
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER if role=='major' else WD_ALIGN_PARAGRAPH.LEFT
        if 'indent_pt' in item:pf.first_line_indent=Pt(float(item['indent_pt']))
        if 'left_indent_pt' in item:pf.left_indent=Pt(float(item['left_indent_pt']))
        if 'runs' in item:
            for run in item['runs']:
                add_text(p,run['text'],run.get('role',role),profile,run.get('bold'),run.get('size_pt'))
        else:add_text(p,item.get('text',''),role,profile)
        paragraph_mark(p)
    d.core_properties.title=data.get('title','')
    d.core_properties.subject=''
    d.core_properties.comments=''
    # Remove obsolete content relations from the new copy, while keeping all
    # styles/themes/settings and relationships actually referenced by cloned content.
    used=set(body.xpath('//@r:embed | //@r:link | //@r:id'))
    for rid,rel in list(d.part.rels.items()):
        if rel.reltype.rsplit('/',1)[-1] in {'image','hyperlink'} and rid not in used:d.part.drop_rel(rid)
    output=Path(output);output.parent.mkdir(parents=True,exist_ok=True);d.save(output)
    return output

if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('input',type=Path);parser.add_argument('output',type=Path)
    parser.add_argument('--profile',choices=list(SOURCES),default='approved')
    args=parser.parse_args()
    print(build(json.loads(args.input.read_text()),args.output,args.profile))
